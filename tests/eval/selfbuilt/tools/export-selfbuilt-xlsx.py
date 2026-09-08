import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


def rows(path):
    return [json.loads(line) for line in Path(path).read_text(encoding="utf-8").splitlines() if line]


def header(ws, labels):
    ws.append(labels)
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def add_table(ws, name):
    table = Table(displayName=name, ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(table)


def widths(ws, mapping):
    for column, width in mapping.items():
        ws.column_dimensions[column].width = width


def main():
    if len(sys.argv) != 4:
        raise SystemExit("usage: export-selfbuilt-xlsx.py cases.jsonl manifest.json output.xlsx")
    cases = rows(sys.argv[1])
    manifest = json.loads(Path(sys.argv[2]).read_text(encoding="utf-8"))
    wb = Workbook()
    info = wb.active
    info.title = "说明"
    info.merge_cells("A1:F1")
    info["A1"] = "Mengshu 自建评测集"
    info["A1"].font = Font(name="Arial", size=18, bold=True, color="FFFFFF")
    info["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    info["A1"].alignment = Alignment(horizontal="left", vertical="center")
    info.row_dimensions[1].height = 28
    details = [
        ("Dataset ID", manifest["datasetId"]),
        ("Version", manifest["datasetVersion"]),
        ("Cases", manifest["caseCount"]),
        ("Dev / Test", f'{manifest["splitCounts"]["dev"]} / {manifest["splitCounts"]["test"]}'),
        ("SHA-256", manifest["casesSha256"]),
        ("Source", manifest["sourceType"]),
        ("Score authority", manifest["scoreAuthority"]),
        ("Formal release eligible", str(manifest["formalReleaseEligible"]).lower()),
        ("说明", "纯确定性合成数据，不读取生产正文或开源 benchmark；SBS 不替代 GMS/PMS。"),
        ("生成", "npm run eval:selfbuilt:prepare"),
        ("评测", "npm run eval:selfbuilt:round1 / round2 / finalize"),
    ]
    for row, (key, value) in enumerate(details, start=3):
        info.cell(row, 1, key).font = Font(name="Arial", bold=True, color="1F4E78")
        info.cell(row, 2, value).font = Font(name="Arial")
    widths(info, {"A": 24, "B": 96})

    case_sheet = wb.create_sheet("Cases")
    case_headers = [
        "id", "split", "language", "dataset_version", "capability", "scenario",
        "expected_mode", "top_k", "query_text", "query_occurred_at", "tenant_id", "user_id",
        "app_id", "workspace_id", "project_id", "agent_id", "namespace", "session_id",
        "visibility", "required_evidence_refs", "forbidden_evidence_refs", "event_count",
    ]
    header(case_sheet, case_headers)
    for case in cases:
        query, gold, scope = case["query"], case["gold"], case["query"]["scope"]
        case_sheet.append([
            case["id"], case["split"], case["language"], case["datasetVersion"],
            case["capability"], case["scenario"], query["expectedMode"], query["topK"],
            query["text"], query["occurredAt"], scope["tenantId"], scope["userId"],
            scope["appId"], scope["workspaceId"], scope["projectId"], scope["agentId"],
            scope["namespace"], scope["sessionId"], scope["visibility"],
            " | ".join(gold["requiredEvidenceRefs"]), " | ".join(gold["forbiddenEvidenceRefs"]),
            len(case["memoryStream"]),
        ])
    add_table(case_sheet, "SelfBuiltCases")
    widths(case_sheet, {
        "A": 25, "B": 9, "C": 10, "D": 14, "E": 36, "F": 23, "G": 15, "H": 8,
        "I": 42, "J": 25, "K": 20, "L": 12, "M": 12, "N": 16, "O": 16,
        "P": 14, "Q": 14, "R": 18, "S": 12, "T": 34, "U": 34, "V": 12,
    })

    event_sheet = wb.create_sheet("Events")
    event_headers = [
        "case_id", "split", "language", "capability", "scenario", "event_id", "evidence_ref",
        "occurred_at", "valid_from", "valid_to", "tenant_id", "user_id", "app_id",
        "workspace_id", "project_id", "agent_id", "namespace", "session_id", "visibility",
        "semantic_type", "kind", "lifecycle_status", "admission_route", "source_class",
        "hydration_state", "superseded_by", "text",
    ]
    header(event_sheet, event_headers)
    for case in cases:
        for event in case["memoryStream"]:
            scope = event["scope"]
            event_sheet.append([
                case["id"], case["split"], case["language"], case["capability"], case["scenario"],
                event["eventId"], event["evidenceRef"], event["occurredAt"], event["validFrom"],
                event.get("validTo", ""), scope["tenantId"], scope["userId"], scope["appId"],
                scope["workspaceId"], scope["projectId"], scope["agentId"], scope["namespace"],
                scope["sessionId"], scope["visibility"], event["semanticType"], event["kind"],
                event["lifecycleStatus"], event["admissionRoute"], event["sourceClass"],
                event.get("hydrationState", "available"), event.get("supersededBy", ""), event["text"],
            ])
    add_table(event_sheet, "SelfBuiltEvents")
    widths(event_sheet, {
        "A": 25, "B": 9, "C": 10, "D": 36, "E": 23, "F": 26, "G": 34, "H": 25,
        "I": 25, "J": 25, "K": 20, "L": 12, "M": 12, "N": 16, "O": 16,
        "P": 14, "Q": 14, "R": 18, "S": 12, "T": 18, "U": 15, "V": 18,
        "W": 18, "X": 16, "Y": 18, "Z": 26, "AA": 54,
    })

    coverage = wb.create_sheet("Coverage")
    header(coverage, ["dimension", "name", "actual", "expected", "status"])
    row = 2
    for name, expected in manifest["countsByCapability"].items():
        coverage.append(["capability", name, f'=COUNTIF(Cases!$E:$E,B{row})', expected,
                         f'=IF(C{row}=D{row},"PASS","CHECK")'])
        row += 1
    for name, expected in manifest["countsByScenario"].items():
        coverage.append(["scenario", name, f'=COUNTIF(Cases!$F:$F,B{row})', expected,
                         f'=IF(C{row}=D{row},"PASS","CHECK")'])
        row += 1
    for name, expected in manifest["splitCounts"].items():
        coverage.append(["split", name, f'=COUNTIF(Cases!$B:$B,B{row})', expected,
                         f'=IF(C{row}=D{row},"PASS","CHECK")'])
        row += 1
    coverage["G1"], coverage["H1"], coverage["I1"], coverage["J1"], coverage["K1"] = (
        "capability", "scenario", "actual", "expected", "status")
    for cell in coverage[1][6:11]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    cross_row = 2
    for capability, scenarios in manifest["crossDistribution"].items():
        for scenario, expected in scenarios.items():
            coverage.cell(cross_row, 7, capability)
            coverage.cell(cross_row, 8, scenario)
            coverage.cell(cross_row, 9, f'=COUNTIFS(Cases!$E:$E,G{cross_row},Cases!$F:$F,H{cross_row})')
            coverage.cell(cross_row, 10, expected)
            coverage.cell(cross_row, 11, f'=IF(I{cross_row}=J{cross_row},"PASS","CHECK")')
            cross_row += 1
    widths(coverage, {"A": 16, "B": 38, "C": 14, "D": 14, "E": 12,
                      "G": 38, "H": 24, "I": 14, "J": 14, "K": 12})
    green = PatternFill("solid", fgColor="E2F0D9")
    coverage.conditional_formatting.add(f"E2:E{row - 1}",
                                        CellIsRule(operator="equal", formula=['"PASS"'], fill=green))
    coverage.conditional_formatting.add(f"K2:K{cross_row - 1}",
                                        CellIsRule(operator="equal", formula=['"PASS"'], fill=green))
    for ws in wb.worksheets:
        for row_cells in ws.iter_rows():
            for cell in row_cells:
                if cell.row != 1:
                    cell.font = Font(name="Arial", bold=cell.font.bold, color=cell.font.color)
                cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (2, 9, 20, 21, 27))
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    Path(sys.argv[3]).parent.mkdir(parents=True, exist_ok=True)
    wb.save(sys.argv[3])


if __name__ == "__main__":
    main()
